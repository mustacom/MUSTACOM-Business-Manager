"""Table exporters: CSV / Excel / PDF used by every report and list screen."""

from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

from PySide6.QtGui import QTextDocument

from ..core.money import cents_to_money


def _cell(columns: list[tuple[str, str, str]], row: dict, key: str, kind: str):
    value = row.get(key)
    if value is None:
        return ""
    if kind == "money":
        return float(cents_to_money(int(value)))
    if kind == "amount":
        return float(value) if isinstance(value, (int, float, Decimal)) else float(Decimal(str(value)))
    if kind in ("qty",):
        return float(Decimal(int(value)) / 1000)
    if kind == "percent":
        return float(value) / 100
    return value


def export_rows(path: Path | str, columns: list[tuple[str, str, str]],
                rows: list[dict]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = [c[1] for c in columns]
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Export"
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([_cell(columns, row, key, kind) for key, _h, kind in columns])
        for index, column in enumerate(columns, start=1):
            width = max(len(str(column[1])) + 2, 12)
            sheet.column_dimensions[get_column_letter(index)].width = width
        workbook.save(path)
    else:
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerow(headers)
            for row in rows:
                writer.writerow([_cell(columns, row, key, kind)
                                 for key, _h, kind in columns])
    return path


def export_table_pdf(title: str, columns: list[tuple[str, str, str]], rows: list[dict],
                     path: Path | str, company_header: str = "") -> Path:
    """Simple tabular PDF report (used by Reports and printable listings)."""
    from .render import doc_css, esc, Paper

    paper = Paper("a4")
    header_rows = "".join(f"<th>{esc(h)}</th>" for _k, h, _t in columns)
    body_rows = []
    for row in rows:
        cells = []
        for key, _h, kind in columns:
            value = row.get(key)
            if kind == "money":
                value = f"{cents_to_money(int(value or 0)):,.2f}".replace(",", " ")
            elif kind == "qty":
                value = f"{Decimal(int(value or 0)) / 1000:,.3f}"
            cells.append(f'<td class="{"num" if kind in ("money","qty","int") else ""}">'
                         f"{esc(value)}</td>")
        body_rows.append("<tr>" + "".join(cells) + "</tr>")
    document = QTextDocument()
    document.setDocumentMargin(0)
    document.setDefaultStyleSheet(doc_css(paper, False))
    document.setHtml(
        f"<div class='company'>{esc(company_header)}</div>"
        f"<h2 style='color:#0F766E'>{esc(title)}</h2>"
        f"<table class='lines'><tr>{header_rows}</tr>{''.join(body_rows)}</table>"
        f"<div class='footer'>{len(rows)} ligne(s)</div>")
    from .render import document_to_pdf

    return document_to_pdf(None, document, path, "a4")
